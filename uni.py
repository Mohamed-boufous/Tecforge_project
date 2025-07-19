import streamlit as st
import json
import math
import weaviate
import queue
import weaviate.classes.query as wq
import weaviate.classes.config as wvc
from sentence_transformers import SentenceTransformer
import os
import requests
import zipfile
import io
import time
from pypdf import PdfReader
import docx
import pandas as pd
import shutil
from PIL import Image
import pytesseract
from pdf2image import convert_from_path
from datetime import datetime, timezone
from pathlib import Path
import concurrent.futures
from pymongo import MongoClient
import certifi
from dotenv import load_dotenv
import subprocess

# --- Imports spécifiques à Windows ---
if os.name == 'nt':
    import win32com.client as win32
    import pythoncom

# --- Configuration de la Page et des Constantes ---
st.set_page_config(layout="wide", page_title="Assistant d'Appels d'Offres")

NOM_DU_MODELE_DE_VECTEUR = 'BAAI/bge-base-en-v1.5'
CLASS_NAME = "DocumentParagraph"
FILES_DIRECTORY = os.path.join(os.path.dirname(__file__), "documents")

# --- Configurez ces chemins selon votre installation ---
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
POPPLER_PATH = r"C:\poppler-24.02.0\Library\bin"

# --- Fonctions Utilitaires et de Chargement ---

@st.cache_resource
def load_model():
    """Charge le modèle de vectorisation une seule fois."""
    return SentenceTransformer(NOM_DU_MODELE_DE_VECTEUR)

@st.cache_data(ttl=3600)
def load_data_from_mongo():
    """Charge les données des appels d'offres directement depuis MongoDB Atlas."""
    MONGO_URI = os.getenv("MONGO_URI")
    if not MONGO_URI:
        st.error("La variable d'environnement MONGO_URI n'est pas définie !")
        return [], {}
    try:
        client = MongoClient(MONGO_URI, tls=True, tlsCAFile=certifi.where(), serverSelectionTimeoutMS=10000)
        db = client.safakate_db
        collection = db.consultations
        data = list(collection.find({}).sort("publishedDate", -1))
        client.close()

        acheteurs, provinces, domaines = set(), set(), set()
        for item in data:
            if item.get("acheteur"): acheteurs.add(item["acheteur"])
            if isinstance(item.get("provinces"), list): provinces.update(item["provinces"])
            if isinstance(item.get("domains"), list):
                for domain_item in item["domains"]:
                    if domain_item.get("domain"): domaines.add(domain_item["domain"])
        
        return data, {
            "acheteurs": sorted(list(acheteurs)),
            "provinces": sorted(list(provinces)),
            "domaines": sorted(list(domaines))
        }
    except Exception as e:
        st.error(f"Erreur de connexion à MongoDB : {e}")
        return [], {}

def format_date(date_string):
    """Formate une date ISO en format lisible."""
    if not date_string: return "N/A"
    try:
        dt_object = datetime.fromisoformat(date_string.replace('Z', '+00:00'))
        return dt_object.strftime("%A %d/%m/%Y %H:%M")
    except (ValueError, TypeError): return "Date invalide"

def jours_restants(date_string):
    """Calcule le nombre de jours restants avant une date."""
    if not date_string: return ""
    try:
        now = datetime.now(timezone.utc)
        end_date = datetime.fromisoformat(date_string.replace('Z', '+00:00'))
        delta = end_date - now
        return f"⏳ Il reste {delta.days} jour(s)" if delta.days >= 0 else "Terminé"
    except (ValueError, TypeError): return ""

# --- Fonctions d'Extraction de Texte ---

def extraire_texte_images_pdf_ocr(chemin_fichier):
    texte_ocr = ""
    try:
        images = convert_from_path(chemin_fichier, poppler_path=POPPLER_PATH)
        for img in images:
            texte_ocr += pytesseract.image_to_string(img, lang='fra') + "\n"
    except Exception as e:
        st.warning(f"Avertissement OCR sur {os.path.basename(chemin_fichier)}: {e}")
    return texte_ocr

def extraire_texte_pdf(chemin_fichier):
    """Retourne le texte ET un booléen indiquant si l'OCR a été nécessaire."""
    texte_normal = ""
    ocr_utilise = False
    try:
        with open(chemin_fichier, "rb") as f:
            reader = PdfReader(f)
            for page in reader.pages:
                contenu = page.extract_text()
                if contenu: texte_normal += contenu + "\n"
    except Exception: pass

    if len(texte_normal.strip()) < 100:
        texte_ocr = extraire_texte_images_pdf_ocr(chemin_fichier)
        if texte_ocr:
            ocr_utilise = True
            return texte_ocr, ocr_utilise
    return texte_normal, ocr_utilise

def extraire_texte_docx(chemin_fichier):
    try:
        document = docx.Document(chemin_fichier)
        return "\n".join([para.text for para in document.paragraphs if para.text.strip()])
    except Exception as e: raise Exception(f"Erreur DOCX: {e}")

def extraire_texte_excel(chemin_fichier):
    try:
        df = pd.read_excel(chemin_fichier, sheet_name=None, header=None)
        texte = ""
        for sheet_name, sheet_df in df.items():
            texte += sheet_df.to_string(index=False, header=False) + "\n"
        return texte
    except Exception as e: raise Exception(f"Erreur Excel: {e}")

def decouper_texte(texte):
    return [p.strip() for p in texte.split("\n") if len(p.strip()) > 10]

# --- Fonctions de Traitement et d'Indexation (Multithreading Corrigé) ---

def traiter_fichier(client, chemin_fichier, model, progress_queue):
    """
    Intérêt : Traite les fichiers par lots et envoie des mises à jour
    de progression via une file d'attente pour un suivi en temps réel.
    """
    nom_fichier = os.path.basename(chemin_fichier)
    try:
        ocr_utilise = False
        extension = os.path.splitext(chemin_fichier)[1].lower()
        texte = ""

        # Étape 1 : Extraction du texte (représente ~10% du travail)
        progress_queue.put((nom_fichier, 5, "Extraction du texte..."))
        if extension == ".pdf":
            texte, ocr_utilise = extraire_texte_pdf(chemin_fichier)
        elif extension == ".docx":
            texte = extraire_texte_docx(chemin_fichier)
        elif extension in [".xlsx", ".xls"]:
            texte = extraire_texte_excel(chemin_fichier)
        
        if not texte:
            progress_queue.put((nom_fichier, 100, "Fichier vide ou illisible"))
            return 0, ocr_utilise

        paragraphes = decouper_texte(texte)
        if not paragraphes:
            progress_queue.put((nom_fichier, 100, "Aucun paragraphe trouvé"))
            return 0, ocr_utilise

        # Étape 2 : Vectorisation et insertion par lots
        total_paragraphes = len(paragraphes)
        batch_size = 32  # Traiter 32 paragraphes à la fois
        doc_collection = client.collections.get(CLASS_NAME)

        for i in range(0, total_paragraphes, batch_size):
            batch_paragraphes = paragraphes[i:i + batch_size]
            
            # Vectorisation du lot
            batch_embeddings = model.encode(batch_paragraphes, show_progress_bar=False)
            
            # Insertion du lot
            objects_to_insert = [
                weaviate.classes.data.DataObject(properties={"content": p, "source": nom_fichier}, vector=emb.tolist())
                for p, emb in zip(batch_paragraphes, batch_embeddings)
            ]
            if objects_to_insert:
                doc_collection.data.insert_many(objects_to_insert)

            # Calcul et envoi de la progression (de 10% à 95%)
            progress_percentage = 10 + int(((i + len(batch_paragraphes)) / total_paragraphes) * 85)
            progress_queue.put((nom_fichier, progress_percentage, f"Traitement... {i + len(batch_paragraphes)}/{total_paragraphes}"))

        # Étape finale
        progress_queue.put((nom_fichier, 100, f"✅ Terminé ({total_paragraphes} paragraphes)"))
        return total_paragraphes, ocr_utilise

    except Exception as e:
        # En cas d'erreur, on l'envoie via la queue
        progress_queue.put((nom_fichier, -1, str(e)))
        return 0, False




def convertir_vers_docx(dossier_path):
    """
    Convertit les .doc et .rtf en .docx en utilisant LibreOffice.
    Cette méthode est fiable et préserve le formatage.
    """
    # 1. On trouve les fichiers à convertir
    extensions = (".doc", ".rtf")
    fichiers_a_convertir = [f for f in os.listdir(dossier_path) if f.lower().endswith(extensions) and not f.startswith('~$')]

    if not fichiers_a_convertir:
        st.info("Aucun fichier .doc ou .rtf à convertir.")
        return

    placeholder = st.empty()
    fichiers_convertis = 0

    for nom_fichier in fichiers_a_convertir:
        chemin_original = os.path.join(dossier_path, nom_fichier)
        placeholder.info(f"🔄 Conversion de {nom_fichier} avec LibreOffice...")

        try:
            # 2. On prépare et exécute la commande LibreOffice
            commande = [
                "soffice",
                "--headless",        # Exécute sans ouvrir de fenêtre
                "--convert-to", "docx", # Format de sortie
                "--outdir", dossier_path, # Dossier de destination
                chemin_original      # Fichier à convertir
            ]
            result = subprocess.run(commande, check=True, capture_output=True, timeout=120)

            # 3. Si la conversion réussit, on supprime l'ancien fichier
            os.remove(chemin_original)
            fichiers_convertis += 1

        except FileNotFoundError:
            st.error("❌ Commande 'soffice' introuvable. Vérifiez que l'Étape 2 (ajout au PATH) a bien été effectuée.")
            return # Inutile de continuer
        except subprocess.CalledProcessError as e:
            st.warning(f"⚠️ La conversion de '{nom_fichier}' a échoué. Erreur : {e.stderr.decode()}")
        except Exception as e:
            st.warning(f"⚠️ Une erreur est survenue lors de la conversion de '{nom_fichier}': {e}")

    placeholder.empty()
    if fichiers_convertis > 0:
        st.success(f"{fichiers_convertis} fichier(s) ont été convertis en .docx.")

def generer_liens(lien_initial: str):
    lien_demande = lien_initial.replace("entreprise.EntrepriseDetailsConsultation", "entreprise.EntrepriseDemandeTelechargementDce")
    lien_final = lien_demande.replace("entreprise.EntrepriseDemandeTelechargementDce", "entreprise.EntrepriseDownloadCompleteDce")
    lien_final = lien_final.replace("refConsultation=", "reference=")
    lien_final = lien_final.replace("orgAcronyme=", "orgAcronym=")
    return lien_final


# Vous pouvez ajouter cette nouvelle fonction d'aide au-dessus de la fonction principale
def extraire_et_aplatir_zip(zip_file_object, destination_folder):
    """
    Extrait un fichier zip, gère les zips imbriqués et place tous les fichiers
    à la racine du dossier de destination, en évitant les doublons de noms.
    """
    # Boucle sur chaque élément du fichier zip
    for member in zip_file_object.infolist():
        if member.is_dir():
            continue # On ignore les dossiers

        file_name = os.path.basename(member.filename)
        if not file_name:
            continue

        # Cible où extraire le fichier
        target_path = os.path.join(destination_folder, file_name)

        # --- Gestion des doublons de noms de fichiers ---
        # Si un fichier du même nom existe déjà, on le renomme (ex: 'CPS (1).pdf')
        counter = 1
        original_target_path = target_path
        while os.path.exists(target_path):
            name, ext = os.path.splitext(original_target_path)
            target_path = f"{name} ({counter}){ext}"
            counter += 1
        # ---------------------------------------------------
        
        source = zip_file_object.open(member)
        
        # --- Gestion des zips imbriqués (récursivité) ---
        if member.filename.lower().endswith('.zip'):
            # Si le fichier est un autre zip, on le lit en mémoire et on relance la fonction
            nested_zip_data = io.BytesIO(source.read())
            with zipfile.ZipFile(nested_zip_data, 'r') as nested_zip_ref:
                extraire_et_aplatir_zip(nested_zip_ref, destination_folder)
        # ------------------------------------------------
        else:
            # Pour tous les autres fichiers, on les écrit dans la destination
            with open(target_path, "wb") as target:
                shutil.copyfileobj(source, target)
def traiter_fichier(client, chemin_fichier, model, progress_queue):
    """
    Intérêt : Traite les fichiers par lots et envoie des mises à jour
    de progression via une file d'attente pour un suivi en temps réel.
    """
    nom_fichier = os.path.basename(chemin_fichier)
    try:
        ocr_utilise = False
        extension = os.path.splitext(chemin_fichier)[1].lower()
        texte = ""

        # Étape 1 : Extraction du texte (représente ~10% du travail)
        progress_queue.put((nom_fichier, 5, "Extraction du texte..."))
        if extension == ".pdf":
            texte, ocr_utilise = extraire_texte_pdf(chemin_fichier)
        elif extension == ".docx":
            texte = extraire_texte_docx(chemin_fichier)
        elif extension in [".xlsx", ".xls"]:
            texte = extraire_texte_excel(chemin_fichier)
        
        if not texte:
            progress_queue.put((nom_fichier, 100, "Fichier vide ou illisible"))
            return 0, ocr_utilise

        paragraphes = decouper_texte(texte)
        if not paragraphes:
            progress_queue.put((nom_fichier, 100, "Aucun paragraphe trouvé"))
            return 0, ocr_utilise

        # Étape 2 : Vectorisation et insertion par lots
        total_paragraphes = len(paragraphes)
        batch_size = 32  # Traiter 32 paragraphes à la fois
        doc_collection = client.collections.get(CLASS_NAME)

        for i in range(0, total_paragraphes, batch_size):
            batch_paragraphes = paragraphes[i:i + batch_size]
            
            # Vectorisation du lot
            batch_embeddings = model.encode(batch_paragraphes, show_progress_bar=False)
            
            # Insertion du lot
            objects_to_insert = [
                weaviate.classes.data.DataObject(properties={"content": p, "source": nom_fichier}, vector=emb.tolist())
                for p, emb in zip(batch_paragraphes, batch_embeddings)
            ]
            if objects_to_insert:
                doc_collection.data.insert_many(objects_to_insert)

            # Calcul et envoi de la progression (de 10% à 95%)
            progress_percentage = 10 + int(((i + len(batch_paragraphes)) / total_paragraphes) * 85)
            progress_queue.put((nom_fichier, progress_percentage, f"Traitement... {i + len(batch_paragraphes)}/{total_paragraphes}"))

        # Étape finale
        progress_queue.put((nom_fichier, 100, f"✅ Terminé ({total_paragraphes} paragraphes)"))
        return total_paragraphes, ocr_utilise

    except Exception as e:
        # En cas d'erreur, on l'envoie via la queue
        progress_queue.put((nom_fichier, -1, str(e)))
        return 0, False

def process_files_threaded(client, model, fichiers_paths):
    """
    Intérêt : Utilise une file d'attente (queue) pour recevoir les mises à jour
    de progression en temps réel depuis les threads et mettre à jour l'interface.
    """
    st.subheader("📊 Progression du Traitement des Fichiers")
    
    progress_queue = queue.Queue()
    total_paragraphes_total = 0

    progress_placeholders = {
        os.path.basename(p): (st.text(f"⏳ En attente: {os.path.basename(p)}"), st.progress(0))
        for p in fichiers_paths
    }

    with concurrent.futures.ThreadPoolExecutor() as executor:
        futures = {
            executor.submit(traiter_fichier, client, path, model, progress_queue): os.path.basename(path)
            for path in fichiers_paths
        }

        tasks_done = 0
        total_tasks = len(fichiers_paths)
        
        while tasks_done < total_tasks:
            try:
                nom_fichier, progress, message = progress_queue.get(timeout=0.1)
                
                status_text, progress_bar = progress_placeholders[nom_fichier]
                
                if progress < 0:
                    status_text.error(f"❌ Erreur sur {nom_fichier}")
                    st.error(f"Détail de l'erreur pour '{nom_fichier}': {message}")
                    progress_bar.empty()
                    tasks_done += 1
                else:
                    # --- MODIFIÉ : Ajout du nom du fichier dans le message de statut ---
                    # Intérêt : Le nom du fichier est maintenant toujours visible pendant la progression.
                    status_text.text(f"{nom_fichier}: {message}")
                    # ------------------------------------------------------------------
                    progress_bar.progress(progress / 100.0)
                    if progress == 100:
                        tasks_done += 1
            
            except queue.Empty:
                continue

        for future in futures:
            try:
                nb_paras, ocr_utilise = future.result()
                if ocr_utilise:
                    st.session_state.ocr_files.add(futures[future])
                total_paragraphes_total += nb_paras
            except Exception:
                pass

    return total_paragraphes_total
# Remplacez votre fonction existante par celle-ci
def telecharger_et_indexer_dossier(lien_initial, client, model):
    with st.status("🚀 Démarrage du processus...", expanded=True) as status:
        try:
            status.update(label="🧹 Nettoyage de la base de données et des anciens fichiers...")
            if client.collections.exists(CLASS_NAME): client.collections.delete(CLASS_NAME)
            client.collections.create(
                name=CLASS_NAME,
                properties=[
                    wvc.Property(name="content", data_type=wvc.DataType.TEXT),
                    wvc.Property(name="source", data_type=wvc.DataType.TEXT)
                ],
                vectorizer_config=wvc.Configure.Vectorizer.none()
            )
            if os.path.exists(FILES_DIRECTORY): shutil.rmtree(FILES_DIRECTORY)
            os.makedirs(FILES_DIRECTORY)
            time.sleep(1)

            status.update(label="📥 Téléchargement du dossier...")
            url_dossier = generer_liens(lien_initial)
            response = requests.get(url_dossier, headers={"User-Agent": "Mozilla/5.0"}, timeout=60)
            response.raise_for_status()
            
            status.update(label="📦 Décompression intelligente des fichiers...")
            with zipfile.ZipFile(io.BytesIO(response.content), 'r') as zip_ref:
                extraire_et_aplatir_zip(zip_ref, FILES_DIRECTORY)
            
            with st.expander("🔄 Fichiers .doc en cours de conversion", expanded=True):
                convertir_vers_docx(FILES_DIRECTORY)
                if st.session_state.get('conversion_files'):
                    st.success(f"{len(st.session_state.conversion_files)} fichier(s) converti(s) en .docx.")
                

            # --- MODIFIÉ : Ajout d'un filtre pour ignorer les fichiers temporaires ---
            # Intérêt : On ne traite que les fichiers valides et on ignore ceux qui commencent par '~$' pour éviter les erreurs.
            extensions_valides = (".pdf", ".docx", ".xlsx", ".xls")
            all_files_in_dir = os.listdir(FILES_DIRECTORY)
            fichiers_a_traiter_paths = [
                os.path.join(FILES_DIRECTORY, f) 
                for f in all_files_in_dir 
                if f.lower().endswith(extensions_valides) and not f.startswith('~$')
            ]
            # -------------------------------------------------------------------------

            if not fichiers_a_traiter_paths:
                status.update(label="⚠️ Aucun fichier compatible trouvé dans le ZIP. Le processus est terminé.", state="complete")
                time.sleep(3)
                st.rerun()
                return

            total_paragraphes = process_files_threaded(client, model, fichiers_a_traiter_paths)
            
            with st.expander("👁️ Fichiers PDF ayant nécessité une lecture OCR", expanded=True):
                if st.session_state.get('ocr_files'):
                    for f in st.session_state.ocr_files: st.write(f"• {f}")
                else: st.info("Aucun PDF n'a nécessité d'OCR.")

            status.update(label=f"🎉 Processus terminé ! {total_paragraphes} paragraphes indexés.", state="complete")
            st.balloons()
            time.sleep(3)
            st.rerun()

        except Exception as e:
            status.update(label=f"❌ Erreur critique : {e}", state="error")


def display_list_view(data, filter_options):
    st.markdown("""
    <style>
    .card:hover { box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1), 0 2px 4px -1px rgba(0, 0, 0, 0.06); }
    .badge-en-cours { background-color:#dcfce7; color:#166534; padding: 5px 12px; border-radius: 9999px; font-weight: 600; text-align: center; font-size: 0.875rem; margin-top: 10px; }
    .lot-title { font-size:1.2rem; font-weight:bold; color:#374151; text-align:center; margin:20px 0; padding:10px; background-color:#f9fafb; border-radius:8px; }
    </style>""", unsafe_allow_html=True)
    st.sidebar.header("🔎 Filtres")
    keyword_filter = st.sidebar.text_input("Rechercher par Réf, ID ou Objet")
    acheteur_filter = st.sidebar.multiselect("Filtrer par Acheteur", options=filter_options["acheteurs"])
    province_filter = st.sidebar.multiselect("Filtrer par Province", options=filter_options["provinces"])
    domaine_filter = st.sidebar.multiselect("Filtrer par Domaine", options=filter_options["domaines"])

    filtered_data = data
    if keyword_filter:
        kw = keyword_filter.lower()
        filtered_data = [item for item in filtered_data if kw in str(item.get("consId", "")).lower() or kw in item.get("reference", "").lower() or any(kw in lot.get("lotObject", "").lower() for lot in item.get("lots", []))]
    if acheteur_filter: filtered_data = [item for item in filtered_data if item.get("acheteur") in acheteur_filter]
    if province_filter: filtered_data = [item for item in filtered_data if any(p in province_filter for p in item.get("provinces", []))]
    if domaine_filter: filtered_data = [item for item in filtered_data if any(d.get("domain") in domaine_filter for d in item.get("domains", []))]

    st.title("📄 Appels d'Offres Publics")
    st.write(f"**{len(filtered_data)}** résultat(s) trouvé(s)")
    st.divider()

    ITEMS_PER_PAGE = 10
    total_pages = math.ceil(len(filtered_data) / ITEMS_PER_PAGE) if filtered_data else 1
    if 'page' not in st.session_state or st.session_state.page > total_pages: st.session_state.page = 1
    start_index = (st.session_state.page - 1) * ITEMS_PER_PAGE
    paginated_data = filtered_data[start_index:start_index + ITEMS_PER_PAGE]

    for item in paginated_data:
        with st.container(border=True):
            col1, col2 = st.columns([0.8, 0.2])
            with col1:
                st.markdown(f'<h5>{item.get("AchAbr", "")} - {item.get("acheteur", "N/A")}</h5>', unsafe_allow_html=True)
                st.caption(f"Publié le : {format_date(item.get('publishedDate'))}")
            with col2: st.markdown('<div class="badge-en-cours">EN COURS</div>', unsafe_allow_html=True)
            st.divider()
            col_info, col_boutons = st.columns([0.7, 0.3])
            with col_info:
                st.markdown(f"""
                <div style="line-height:1.8; font-size:0.95rem;">
                    <span>📋 {item.get("procedureType", "N/A")}</span><br>
                    <span><strong>Référence :</strong> {item.get("reference", "N/A")}</span><br>
                    <span>📍 {', '.join(item.get("provinces", []))}</span><br>
                    <span><strong>Date limite :</strong> {format_date(item.get("endDate"))}</span><br>
                    <strong style="color: #d9480f;">{jours_restants(item.get('endDate'))}</strong>
                </div>""", unsafe_allow_html=True)
            with col_boutons:
                st.link_button("🔗 Page de Consultation", item.get("detailsUrl", "#"), use_container_width=True)
                if st.button("⚙️ Traiter ce Dossier", key=f"process_{item.get('consId')}", use_container_width=True):
                    st.session_state.view = 'process'
                    st.session_state.lien_a_traiter = item.get("detailsUrl")
                    st.rerun()
            for lot in item.get("lots", []):
                st.markdown(f'<div class="lot-title">{lot.get("lotObject", "Non spécifié")}</div>', unsafe_allow_html=True)
                c1, c2, c3 = st.columns(3)
                c1.metric("Catégorie", lot.get("lotCategory", "N/A"))
                c2.metric("Estimation", f"{lot.get('lotEstimation', 0):,.2f} MAD".replace(",", " "))
                c3.metric("Caution", f"{lot.get('lotCaution', 0):,.2f} MAD".replace(",", " "))
        st.write("") 

    if total_pages > 1:
        st.divider()
        col_prev, col_input, col_next = st.columns([3, 1, 3])
        if col_prev.button("⬅️ Précédent", disabled=(st.session_state.page <= 1), use_container_width=True):
            st.session_state.page -= 1; st.rerun()
        page_input = col_input.number_input("Page", min_value=1, max_value=total_pages, value=st.session_state.page, key="page_input", label_visibility="collapsed")
        if page_input != st.session_state.page:
            st.session_state.page = page_input; st.rerun()
        if col_next.button("Suivant ➡️", disabled=(st.session_state.page >= total_pages), use_container_width=True):
            st.session_state.page += 1; st.rerun()
        _, col_text_total, _ = st.columns([3, 1, 3])
        col_text_total.markdown(f"<div style='text-align:center;'>sur {total_pages}</div>", unsafe_allow_html=True)

def display_process_view(client, model):
    st.title("⚙️ Traitement et Indexation d'un Appel d'Offres")
    if st.button("⬅️ Retour à la liste"):
        st.session_state.view = 'list'
        if 'lien_a_traiter' in st.session_state: del st.session_state.lien_a_traiter
        st.rerun()
    lien = st.session_state.get("lien_a_traiter", "")
    st.text_input("Lien de la consultation à traiter :", value=lien, disabled=True)
    
    if st.button("Lancer le Traitement", type="primary"):
        if lien:
            st.session_state.conversion_files = []
            st.session_state.ocr_files = set()
            telecharger_et_indexer_dossier(lien, client, model)
        else: st.error("Aucun lien à traiter.")
            
    st.divider()
    st.header("📊 État de la base de données")
    col1, col2 = st.columns(2)
    fichiers_locaux = 0
    if os.path.exists(FILES_DIRECTORY):
        fichiers_locaux = len([f for f in os.listdir(FILES_DIRECTORY) if f.lower().endswith((".pdf", ".docx", ".xlsx", ".xls"))])
    col1.metric(label="📄 Fichiers Locaux Prêts", value=fichiers_locaux)
    
    total_paragraphs = 0
    if client.collections.exists(CLASS_NAME):
        doc_collection = client.collections.get(CLASS_NAME)
        response = doc_collection.aggregate.over_all(total_count=True)
        total_paragraphs = response.total_count
    col2.metric(label="✍️ Paragraphes dans Weaviate", value=total_paragraphs)
    
    st.divider()
    st.header("🔎 Rechercher dans les documents")
    requete_utilisateur = st.text_input("Que cherchez-vous ?", "Fourniture de bureau")
    if st.button("Lancer la recherche"):
        if requete_utilisateur and total_paragraphs > 0:
            vecteur_requete = model.encode(requete_utilisateur).tolist()
            doc_collection = client.collections.get(CLASS_NAME)
            response = doc_collection.query.near_vector(near_vector=vecteur_requete, limit=5, return_metadata=wq.MetadataQuery(distance=True))
            st.subheader("Résultats de la recherche :")
            if not response.objects: st.warning("Aucun résultat trouvé.")
            else:
                for item in response.objects:
                    st.info(f"**Pertinence (distance) :** {item.metadata.distance:.4f} (plus c'est bas, mieux c'est)")
                    st.write(f"📄 **Source** : {item.properties.get('source', 'Inconnue')}")
                    st.write(f"📌 **Paragraphe** : {item.properties.get('content', '')}")
                    st.divider()
        elif not requete_utilisateur: st.warning("Veuillez entrer une requête de recherche.")
        else: st.warning("La base de données est vide. Veuillez d'abord traiter un dossier.")
# AJOUTEZ CETTE NOUVELLE FONCTION

# --- Exécution Principale ---
if 'view' not in st.session_state: st.session_state.view = 'list'
if 'page' not in st.session_state: st.session_state.page = 1

load_dotenv()
data, filter_options = load_data_from_mongo()
model = load_model()

try:
    with weaviate.connect_to_local(port=8080, grpc_port=50051) as client:
        if st.session_state.view == 'list':
            display_list_view(data, filter_options)
        elif st.session_state.view == 'process':
            display_process_view(client, model)
except Exception as e:
    st.error(f"Erreur critique de connexion à Weaviate : {e}")
    st.info("Veuillez vous assurer que votre instance Weaviate est bien en cours d'exécution sur le port 8080.")